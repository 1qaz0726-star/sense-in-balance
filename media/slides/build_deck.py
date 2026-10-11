# 提案簡報產生器：沿用胡馨 NightSync 版的視覺（深藍底、青綠強調、章節小標、黃色註記），內容改為「該說才說」
from pptx import Presentation
from pptx.util import Emu, Pt, Inches
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
import copy, os

IMG = r'F:\tmp\slides\img'
AS = r'F:\tmp\slides\assets'
SHOT = r'C:\Users\葉士愷\Documents\友達智慧座艙\★給隊友_座艙APK\實機截圖_v1.1'
FONT = 'Microsoft JhengHei'
BG = RGBColor(0x0A, 0x13, 0x20); ACC = RGBColor(0x61, 0xDD, 0xD0); HI = RGBColor(0xE7, 0xEB, 0xEC)
BODY = RGBColor(0xC9, 0xD1, 0xD6); LO = RGBColor(0x8A, 0x97, 0xA0); TH = RGBColor(0x20, 0x3A, 0x48)
CARD = RGBColor(0x10, 0x20, 0x2E); LINE = RGBColor(0x2A, 0x46, 0x56); NOTE = RGBColor(0xF2, 0xB5, 0x44)
RED = RGBColor(0xE5, 0x48, 0x4D); ROW = RGBColor(0x0D, 0x18, 0x26)

prs = Presentation()
prs.slide_width = Inches(13.333); prs.slide_height = Inches(7.5)
W, H = 13.333, 7.5
BRAND = '該說才說'

def font_run(r, size, color=BODY, bold=False):
    r.font.name = FONT; r.font.size = Pt(size); r.font.bold = bold; r.font.color.rgb = color
    rPr = r._r.get_or_add_rPr()
    for tag in ('a:ea', 'a:cs'):
        el = rPr.find(qn(tag))
        if el is None:
            el = rPr.makeelement(qn(tag), {}); rPr.append(el)
        el.set('typeface', FONT)

def text(s, x, y, w, h, t, size=18, color=BODY, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, spacing=1.15):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.02); tf.margin_top = tf.margin_bottom = Inches(0.02)
    lines = t if isinstance(t, list) else [t]
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align; p.line_spacing = spacing
        segs = ln if isinstance(ln, list) else [(ln, color, bold)]
        for seg in segs:
            if isinstance(seg, str): seg = (seg, color, bold)
            r = p.add_run(); r.text = seg[0]; font_run(r, size, seg[1], seg[2])
    return tb

def rect(s, x, y, w, h, fill=CARD, line=LINE, lw=1.25, shape=MSO_SHAPE.RECTANGLE, dash=False):
    sh = s.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is None: sh.fill.background()
    else: sh.fill.solid(); sh.fill.fore_color.rgb = fill
    if line is None: sh.line.fill.background()
    else:
        sh.line.color.rgb = line; sh.line.width = Pt(lw)
        if dash:
            from pptx.enum.dml import MSO_LINE
            sh.line.dash_style = MSO_LINE.DASH
    sh.shadow.inherit = False
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE: sh.adjustments[0] = 0.08
    return sh

def boxtext(s, x, y, w, h, t, size=16, color=HI, bold=True, fill=CARD, line=ACC, align=PP_ALIGN.LEFT, sub=None, subsize=13, dash=False):
    rect(s, x, y, w, h, fill, line, dash=dash)
    lines = [[(t, color, bold)]]
    if sub: lines += [[(ln, LO, False)] for ln in sub.split('\n')]
    tb = text(s, x + 0.15, y + 0.08, w - 0.3, h - 0.16, lines, size, color, bold, align, MSO_ANCHOR.MIDDLE)
    for p in tb.text_frame.paragraphs[1:]:
        for r in p.runs: r.font.size = Pt(subsize)
    return tb

def arrow(s, x1, y1, x2, y2, color=LO):
    c = s.shapes.add_connector(1, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = color; c.line.width = Pt(2)
    ln = c.line._get_or_add_ln()
    tail = ln.makeelement(qn('a:tailEnd'), {'type': 'triangle', 'w': 'med', 'len': 'med'}); ln.append(tail)
    return c

def pic(s, path, x, y, w=None, h=None, crop=None):
    p = s.shapes.add_picture(path, Inches(x), Inches(y), Inches(w) if w else None, Inches(h) if h else None)
    if crop:
        p.crop_left, p.crop_top, p.crop_right, p.crop_bottom = crop
    p.line.color.rgb = LINE; p.line.width = Pt(1)
    return p

def table(s, x, y, w, rows, cols, size=14, rowh=0.46, head_color=ACC, first_bold=False, colors=None):
    nr, nc = len(rows), len(rows[0])
    gt = s.shapes.add_table(nr, nc, Inches(x), Inches(y), Inches(w), Inches(rowh * nr))
    tbl = gt.table
    tblPr = tbl._tbl.tblPr
    for k in ('bandRow', 'firstRow'):
        tblPr.set(k, '0')
    style = tblPr.find(qn('a:tableStyleId'))
    if style is not None: style.text = '{5940675A-B579-460E-94D1-54222C63F5DA}'  # No Style, Table Grid
    for j, cw in enumerate(cols): tbl.columns[j].width = Inches(cw)
    for i, row in enumerate(rows):
        tbl.rows[i].height = Inches(rowh)
        for j, val in enumerate(row):
            c = tbl.cell(i, j)
            c.fill.solid(); c.fill.fore_color.rgb = TH if i == 0 else ROW
            c.margin_left = c.margin_right = Inches(0.1); c.margin_top = c.margin_bottom = Inches(0.05)
            c.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf = c.text_frame; tf.word_wrap = True
            p = tf.paragraphs[0]; p.text = ''
            col = head_color if i == 0 else (colors[i][j] if colors and colors[i] and colors[i][j] else (HI if (first_bold and j == 0) else BODY))
            r = p.add_run(); r.text = str(val); font_run(r, size, col, i == 0 or (first_bold and j == 0))
            # 邊框
            tcPr = c._tc.get_or_add_tcPr()
            for side in ('a:lnL', 'a:lnR', 'a:lnT', 'a:lnB'):
                ln = tcPr.makeelement(qn(side), {'w': '9525'})
                sf = ln.makeelement(qn('a:solidFill'), {}); clr = sf.makeelement(qn('a:srgbClr'), {'val': '1E2E3C'}); sf.append(clr); ln.append(sf)
                tcPr.append(ln)
    return gt

def slide(tag=None, title=None, sub=None, section='方案設計', page=None):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid(); s.background.fill.fore_color.rgb = BG
    if tag: text(s, 0.7, 0.42, 8, 0.4, tag, 14, ACC)
    if title: text(s, 0.7, 0.85, 12, 0.9, title, 32, HI, True)
    if sub: text(s, 0.7, 1.68, 12, 0.5, sub, 16, LO)
    if page:
        text(s, 0.7, 7.0, 6, 0.3, f'{BRAND}  /  {section}', 11, LO)
        text(s, 12.0, 7.0, 0.7, 0.3, f'{page:02d}', 12, ACC, align=PP_ALIGN.RIGHT)
    return s

def note(s, t, y=6.55):
    text(s, 0.7, y, 12, 0.35, t, 12.5, NOTE)

P = lambda n: n  # 頁碼

# ── 01 封面 ───────────────────────────────────────────
s = prs.slides.add_slide(prs.slide_layouts[6])
s.background.fill.solid(); s.background.fill.fore_color.rgb = BG
s.shapes.add_picture(AS + r'\cover.jpg', 0, 0, prs.slide_width, prs.slide_height)
text(s, 0.75, 1.55, 7, 1.3, '該說才說', 60, HI, True)
text(s, 0.78, 2.85, 7, 0.7, '會自己動手的安靜座艙 AI', 28, HI, True)
text(s, 0.8, 3.75, 4.6, 1.3, ['系統能自己做的，就自己做；', '只有在有代價、或有危險的時候，', '才用一句話告訴你。'], 16, ACC)
text(s, 0.8, 6.35, 9, 0.35, '2026 AI 座艙感知與視覺化 Hackathon｜實作組', 13, LO)
text(s, 0.8, 6.7, 9, 0.35, '團隊名稱：感知有度 Sense in Balance', 13, LO)

# ── 02 目次 ───────────────────────────────────────────
s = slide('提案總覽', '不是讓車更會說話，而是讓車知道什麼時候不該說', '目次', page=2)
chap = [('01', '團隊背景', '03'), ('02', '提案概述', '04'), ('03', '目標對象與痛點分析', '06'), ('04', '解決方案設計', '07'),
        ('05', 'AI 應用方法', '11'), ('06', '獨特優勢與差異化', '14'), ('07', '預期效益與落地評估', '16')]
for i, (n, t, pg) in enumerate(chap):
    cx = 0.7 + (i % 2) * 6.1; cy = 2.45 + (i // 2) * 0.95
    text(s, cx, cy, 0.9, 0.6, n, 26, ACC, True)
    text(s, cx + 0.95, cy + 0.08, 4.2, 0.5, t, 20, HI, True)
    text(s, cx + 5.0, cy + 0.12, 0.7, 0.5, pg, 14, LO, align=PP_ALIGN.RIGHT)
    rect(s, cx, cy + 0.7, 5.7, 0.012, LINE, None)
note(s, '所有實機畫面皆取自報名影片（Samsung Galaxy S21＋iPad，實際運作錄影）；方案說明影片連結見報名信件。')

# ── 03 團隊背景 ───────────────────────────────────────
s = slide('① 團隊背景', '從感知、AI 到座艙呈現，三個人一條龍', '工程實作＋資訊設計：每個畫面都是真的在跑的 App', page=3)
team = [('葉士愷', '國立高雄大學 土環系', '前後端網頁部署、API 與資料庫應用', '標註資料・程式開發\nAI Box・端側模型・HMI App'),
        ('黃昱羲', '國立高雄大學 土環系', '架網站和 app、商業分析', '感知處理流程・Android HMI'),
        ('胡馨', '國立成功大學 藝術所', '資訊處理、專案管理、藝術行政', '提案簡報・驗證證據整理')]
for i, (n, sc, sk, role) in enumerate(team):
    cx = 0.9 + i * 4.1
    ring = rect(s, cx + 0.95, 2.35, 1.9, 1.9, None, ACC, 2.5, MSO_SHAPE.OVAL)
    s.shapes.add_picture(AS + f'\\avatar_{i}.png', Inches(cx + 1.05), Inches(2.45), Inches(1.7), Inches(1.7))
    text(s, cx, 4.4, 3.8, 0.5, n, 22, HI, True, PP_ALIGN.CENTER)
    text(s, cx, 4.9, 3.8, 0.9, [f'學校：{sc}', f'專長：{sk}'], 12.5, LO, align=PP_ALIGN.CENTER)
    text(s, cx, 5.75, 3.8, 0.7, role.split('\n'), 13.5, ACC, True, PP_ALIGN.CENTER)
note(s, '團隊優勢：工程端能把模型部署到手機上實際執行，設計端負責敘事與證據整理；提案中的每一個畫面都來自實機錄影。', 6.62)

# ── 04 提案概述 1/2 ───────────────────────────────────
s = slide('② 提案概述', '系統先動手，有代價或危險才開口', '設計理念：在 12,000 種座艙狀態組合中，規則判定有 70% 的正確反應是「不說話」', page=4)
pillars = [('01', '自己動手', '能自動處理的事直接做，不叫駕駛動手', '例：車內 CO2 上升 → 系統自動切外循環'),
           ('02', '該說才說', '只有「有代價」才告知，而且只說一句', '例：車外 34 度 →「已換外循環，會稍熱」'),
           ('03', '危險不等模型', '安全警示由規則直接發出，八字以內', '例：先打方向燈 → 「右後方有機車」')]
for i, (n, t, d, e) in enumerate(pillars):
    cx = 0.7 + i * 4.1
    text(s, cx, 2.55, 3.8, 0.7, n, 34, ACC, True)
    text(s, cx, 3.35, 3.8, 0.5, t, 22, HI, True)
    text(s, cx, 3.95, 3.7, 0.9, d, 15, BODY)
    rect(s, cx, 4.85, 3.75, 0.85, CARD, LINE)
    text(s, cx + 0.15, 4.92, 3.5, 0.75, e, 13, LO, anchor=MSO_ANCHOR.MIDDLE)
text(s, 0.7, 5.95, 12, 0.45, [[('核心 AI 全部在車上離線執行：', HI, True), ('規則引擎決定做什麼，端側小模型 Qwen2.5-1.5B 只負責講成一句話，本地知識庫回答停車提問', BODY, False)]], 15)
note(s, '70%：研究原型 states.py 窮舉 12,000 組狀態（CO2、溫度、循環、乘員、車速）之決策分布；App 規則引擎與之逐筆一致。')

# ── 05 提案概述 2/2 ───────────────────────────────────
s = slide('② 提案概述', '四個情境，涵蓋四種緊急度', '一個事件、兩塊螢幕、一個聲音：每一種緊急度都有固定的呈現方式', page=5)
rows = [['情境', '座艙發生什麼', '系統做什麼', '駕駛看到／聽到'],
        ['A 一切正常', '各項數值正常', '不說話（緊急度：不說話）', '儀表只有時速、檔位、油量'],
        ['B 自動換氣', '車內 CO2 慢慢上升', '自動切外循環，告知代價（主動告知 ≤15 字）', '中控卡片＋語音「已換外循環，會稍熱」'],
        ['C 盲點警示', '打方向燈，後方有來車', '規則直出，不等模型（安全 ≤8 字）', '儀表該側紅邊＋來車圖示，語音「右後方有機車」'],
        ['D 停車問答', '停車（P 檔）後提問', '查車上的知識庫（回答 ≤30 字）', '中控問答卡＋語音；查不到就說查不到']]
table(s, 0.7, 2.35, 11.9, rows, [1.7, 2.6, 3.6, 4.0], 14, 0.62, first_bold=True)
note(s, '盲點與 CO2 為模擬注入的事件（Demo 用），事件格式與真實感測一致，決賽後可替換為 CAN／感測器資料。', 5.95)

# ── 06 目標對象與痛點 ─────────────────────────────────
s = slide('③ 目標對象與痛點分析', '車越來越愛說話，但很多事駕駛不需要知道', '目標使用者：開車中的駕駛｜導入對象：車廠與座艙系統整合商（CDC 平台）', page=6)
pains = [('提示過載', '提醒、建議、通知一個接一個，每一則都要駕駛分心看或聽', '→ 預設不說話，該說才說'),
         ('把系統能做的事丟給駕駛', '「請切換外循環」這類提醒，其實系統自己就能做', '→ 執行層：系統先做，再告知代價'),
         ('依賴雲端、可能亂答', '斷網就失效；生成式回答可能編造數字或功能', '→ 核心離線，查不到就說查不到')]
for i, (t, d, a) in enumerate(pains):
    cx = 0.7 + i * 4.1
    rect(s, cx, 2.45, 3.8, 3.3, CARD, LINE)
    text(s, cx + 0.25, 2.65, 3.4, 0.5, f'痛點 {i+1}', 13, ACC)
    text(s, cx + 0.25, 3.05, 3.4, 0.6, t, 20, HI, True)
    text(s, cx + 0.25, 3.75, 3.35, 1.1, d, 14.5, BODY)
    text(s, cx + 0.25, 4.85, 3.35, 0.8, a, 14, ACC, True)
note(s, '以上痛點為團隊依座艙使用情境歸納的設計假設，決賽前補充駕駛訪談驗證；本頁不引用未經查證的統計數字。', 6.0)

# ── 07 解決方案：系統架構 ─────────────────────────────
s = slide('④ 解決方案設計', '系統架構：一個 AI Box，兩塊螢幕', 'AI Box 決定與措辭，透過 WebSocket 每秒推送事件；兩塊螢幕各自依緊急度呈現', page=7)
boxtext(s, 0.7, 3.0, 2.3, 1.7, '車輛訊號', sub='CO2・車溫・車速\n方向燈・盲點（模擬）\n停車提問', size=17)
arrow(s, 3.05, 3.85, 3.55, 3.85)
rect(s, 3.6, 2.3, 5.2, 3.6, None, RGBColor(0x4A, 0x66, 0x76), 1.5)
text(s, 3.8, 2.38, 5, 0.4, 'AI Box（車上・離線）', 14, LO, True)
boxtext(s, 3.85, 2.85, 4.7, 0.62, '規則引擎：決定做什麼、多緊急、有沒有代價', size=14)
boxtext(s, 3.85, 3.6, 4.7, 0.62, '端側小模型 Qwen2.5-1.5B：只講成一句話', size=14)
boxtext(s, 3.85, 4.35, 4.7, 0.62, '檢查器：字數、繁體、不編數字（不過→模板）', size=14)
boxtext(s, 3.85, 5.1, 2.25, 0.62, '本地知識庫', size=14)
boxtext(s, 6.3, 5.1, 2.25, 0.62, '語音（TTS）', size=14)
arrow(s, 8.85, 3.3, 9.6, 2.85); arrow(s, 8.85, 4.6, 9.6, 5.0)
text(s, 8.85, 3.75, 0.9, 0.5, 'WebSocket\n1 秒心跳', 10, LO)
boxtext(s, 9.65, 2.25, 3.0, 1.35, '數位儀表（駕駛正前方）', sub='時速・檔位・油量\n盲點圖示與紅邊', size=15)
boxtext(s, 9.65, 4.45, 3.0, 1.35, '中控（觸控）', sub='告知卡片・問答\n地圖・停車提問', size=15)
boxtext(s, 3.6, 6.05, 5.2, 0.42, '雲端增強（決賽前完成；有網路才用，見 p.12）', size=12, color=LO, bold=False, line=LO, fill=None, align=PP_ALIGN.CENTER, dash=True)
note(s, '儀表超過 3 秒收不到事件即顯示失效燈；安全警示永不經過雲端與模型。Demo 配置：手機＝儀表＋AI Box，iPad 網頁＝中控。', 6.6)

# ── 08 多螢分工 ───────────────────────────────────────
s = slide('④ 解決方案設計', '多螢協作：同一個事件，兩塊螢幕各司其職', '依緊急度決定「誰顯示、顯示多少、要不要出聲」', page=8)
rows = [['緊急度', '數位儀表', '中控', '語音'],
        ['安全（≤8 字）', '該側紅邊＋來車圖示＋方向燈', '紅色警示卡', '立即插話'],
        ['主動告知（≤15 字）', '不顯示（不搶駕駛視線）', '卡片：一句話＋原因＋代價', '念一句'],
        ['回答（≤30 字，限停車）', '—', '問答卡（查不到就說查不到）', '念答案'],
        ['不說話', '只有時速、檔位、油量', '地圖', '—']]
table(s, 0.7, 2.3, 7.4, rows, [1.9, 2.2, 2.1, 1.2], 13, 0.6, first_bold=True)
pic(s, IMG + r'\dash_C.jpg', 8.4, 2.3, 4.25, 2.39)
text(s, 8.4, 4.72, 4.25, 0.3, '儀表：右後方機車（紅邊與圖示都在右側）', 11, LO)
pic(s, IMG + r'\close_B.jpg', 8.4, 5.05, 2.65, 1.49)
text(s, 11.15, 5.1, 1.5, 1.4, ['中控：', '自動換氣', '卡片'], 11, LO)

# ── 09 操作流程（實機） ────────────────────────────────
s = slide('④ 解決方案設計', '操作流程：實機畫面', '手機（儀表＋AI Box）與 iPad（中控）實際運作錄影截圖', page=9)
shots = [('wide_A.jpg', 'A 一切正常：兩塊螢幕都安靜'), ('wide_B.jpg', 'B 自動換氣：中控卡片＋語音'),
         ('wide_C.jpg', 'C 盲點：儀表紅邊＋中控紅卡'), ('kb_D.jpg', 'D 停車問答：答案來自車上知識庫')]
for i, (f, c) in enumerate(shots):
    cx = 0.7 + (i % 2) * 6.05; cy = 2.2 + (i // 2) * 2.2
    pic(s, os.path.join(IMG, f), cx, cy, 3.5, 1.97)
    t1, t2 = c.split('：')
    text(s, cx + 3.7, cy + 0.45, 2.2, 1.2, [[(t1, ACC, True)], [(t2, BODY, False)]], 14)
note(s, '畫面中面紙包已做馬賽克；盲點事件為模擬注入，尚未接攝影機。', 6.65)

# ── 10 回應企業命題 ───────────────────────────────────
s = slide('④ 解決方案設計', '如何回應企業命題需求', '對照簡章「實作組」建議架構與評審鼓勵方向', page=10)
rows = [['命題要求', '我們的做法', '狀態'],
        ['螢幕協作（Multi-display Experience）', '儀表與中控依緊急度分工，同一事件不重複呈現', '已完成'],
        ['AI 感知資訊於不同 Display 的呈現', '統一事件信封（INTERFACE v1.0），各屏各自決定呈現', '已完成'],
        ['端到端情境：座艙資訊→AI→HMI 回饋', '四個情境皆有實機錄影；含語音回饋', '已完成'],
        ['Android HMI App（Android 14+、不依賴品牌 SDK）', 'Flutter 開發，只用標準 Android API；S21（Android 15）實測', '已完成'],
        ['AI Box ↔ HMI App', 'AI Box 以 WebSocket 推送事件（Demo 用 Wi-Fi，目標車上 Ethernet）', '已完成'],
        ['Hybrid AI（Offline First、Cloud Enhanced）', '離線核心已可運作；雲端增強與斷網／復網展示決賽前完成', '進行中'],
        ['可移植 AAOS／CDC', '純標準元件、事件格式獨立於螢幕', '設計原則']]
cl = [None] + [[None, None, ACC]] * 5 + [[None, None, NOTE], [None, None, LO]]
table(s, 0.7, 2.25, 11.9, rows, [4.2, 6.3, 1.4], 13, 0.5, first_bold=True, colors=cl)

# ── 11 AI 應用方法 ────────────────────────────────────
s = slide('⑤ AI 應用方法', '判斷交給規則，說話交給車上的小模型', '模型不決定內容，只把規則的決定講成一句話；講不好就換成固定模板', page=11)
steps = [('座艙狀態', 'CO2、溫度、循環\n車速、乘員'), ('規則引擎', '做什麼・多緊急\n有沒有代價'), ('端側 SLM', 'Qwen2.5-1.5B\n手機 GPU 執行'),
         ('檢查器', '字數・繁體\n不編數字・不誇大'), ('輸出', '畫面＋語音\n不過→模板保底')]
for i, (t, d) in enumerate(steps):
    cx = 0.7 + i * 2.5
    boxtext(s, cx, 2.45, 2.05, 1.45, t, sub=d, size=17, subsize=12.5, align=PP_ALIGN.CENTER)
    if i < 4: arrow(s, cx + 2.08, 3.17, cx + 2.47, 3.17)
rows = [['緊急度', '字數上限', '誰來措辭', '例句（實機）'],
        ['安全', '8 字', '規則直出（不經模型）', '右後方有機車'],
        ['主動告知', '15 字', '端側模型，檢查不過改用模板', '已經自動切換外循環，會稍微變熱'],
        ['回答', '30 字', '本地知識庫', '車內二氧化碳累積會讓人昏沉，定時開外循環換氣'],
        ['不說話', '0', '—', '（不出聲、不跳卡片）']]
table(s, 0.7, 4.25, 11.9, rows, [1.6, 1.4, 3.6, 5.3], 13, 0.45, first_bold=True)
note(s, '實機紀錄：模型曾輸出 26 字（超過 15 字上限），檢查器擋下並改用模板「已換外循環，會稍熱」，畫面同時標示「模板」。', 6.6)

# ── 12 Hybrid：離線 vs 連網 ───────────────────────────
s = slide('⑤ AI 應用方法', 'Hybrid AI：斷網照常運作，連網後更聰明', '依簡章要求分別列出兩種模式的功能；決賽現場展示斷網與恢復連網', page=12)
rows = [['功能', '離線（已完成，影片實測）', '連網增強（決賽前完成）'],
        ['決策', '規則引擎（與研究原型 12,000 組逐筆一致）', '—（決策永遠在本地）'],
        ['措辭', '端側 Qwen2.5-1.5B，手機 GPU，實測 4.6–6.8 秒', '—'],
        ['提問理解', '本地比對 9 題知識庫', 'Jev（TypeSafe）判斷問題該交給誰，附校準信心度'],
        ['即時資訊', '—（明講查不到）', 'Gemini 3.1 Flash Lite＋網路搜尋（旅遊、營業時間等），附來源'],
        ['車輛功能問題', '本地知識庫', '不上網找答案（避免答錯車型）'],
        ['安全警示', '規則直出', '永不經雲端'],
        ['語音', 'Android 內建 TTS（zh-TW）', '—']]
table(s, 0.7, 2.25, 11.9, rows, [2.0, 4.9, 5.0], 13, 0.47, first_bold=True)
note(s, '雲端回答同樣經過檢查器；逾時（Jev 1 秒、雲端 5 秒）或斷網即退回本地。雲端模型與費用估算見團隊評估文件，尚未上線。', 6.2)

# ── 13 AI 價值與實測 ──────────────────────────────────
s = slide('⑤ AI 應用方法', 'AI 帶來的價值：說得少、說得準、斷網也在', '以下數字皆為實測或可重現的測試結果', page=13)
stats = [('70%', '的狀態組合', '規則判定「不說話」\n（12,000 組窮舉）'), ('4.6–6.8 秒', '端側措辭', 'Galaxy S21 GPU\n實機錄影量測'),
         ('離線', '斷網仍可措辭與問答', '影片中關閉行動數據\n後實測'), ('16 項', '自動化測試全數通過', '規則逐筆對照、檢查器\n雙屏連線、網頁中控')]
for i, (big, a, b) in enumerate(stats):
    cx = 0.7 + i * 3.05
    rect(s, cx, 2.4, 2.8, 2.9, CARD, LINE)
    text(s, cx + 0.2, 2.6, 2.4, 0.9, big, 30 if len(big) < 6 else 24, ACC, True)
    text(s, cx + 0.2, 3.55, 2.4, 0.5, a, 16, HI, True)
    text(s, cx + 0.2, 4.1, 2.45, 1.0, b.split('\n'), 12.5, LO)
text(s, 0.7, 5.6, 12, 0.8, [[('價值：', HI, True), ('駕駛少被打斷；需要知道的事一句說清楚；安全警示不受模型延遲影響；車在沒有訊號的地方一樣能用。', BODY, False)]], 15)
note(s, '措辭延遲（4.6–6.8 秒）只影響「主動告知」，安全警示不經模型；延遲期間卡片顯示「端側模型措辭中」。', 6.55)

# ── 14 差異化 ─────────────────────────────────────────
s = slide('⑥ 獨特優勢與差異化', '差別不在模型多大，而在什麼時候開口', '與常見車用語音助理做法對照（團隊歸納，非特定產品）', page=14)
rows = [['面向', '常見做法', '該說才說'],
        ['什麼時候說話', '使用者叫才回應，或事件逐一通知', '預設不說；只有代價或危險才說'],
        ['誰動手', '提醒駕駛去做', '系統先做，再告知代價'],
        ['安全警示', '與一般訊息走同一條管道', '規則直出，不等模型、不靠網路'],
        ['斷網', '主要功能依賴雲端', '核心離線可用，雲端只是增強'],
        ['回答可信度', '模型直接生成', '規則決定內容、模型只措辭、檢查器把關'],
        ['多螢', '同一則訊息多屏重複', '依緊急度分工，儀表只留駕駛需要的']]
table(s, 0.7, 2.25, 11.9, rows, [2.4, 4.4, 5.1], 14, 0.53, first_bold=True, colors=[None] + [[None, LO, ACC]] * 6)
note(s, '「規則決策、模型措辭」本身是已知架構；我們的差異在「沉默為預設＋系統先動手＋安全不經模型」的組合，以及可逐筆驗證。', 6.25)

# ── 15 優勢 ───────────────────────────────────────────
s = slide('⑥ 獨特優勢與差異化', '三個可以當場驗證的優勢', '評審可以在決賽現場逐項檢查', page=15)
adv = [('可驗證', ['規則引擎與研究原型 12,000 組逐筆一致', '每句話標示來源：模型／模板／規則／知識庫', '影片聲音逐段以語音辨識對時驗證']),
       ('可移植', ['純 Flutter／Android 標準 API，不依賴品牌 SDK', '事件格式（INTERFACE v1.0）與螢幕無關', '儀表、中控可換成 8:3 與 16:9 實車螢幕']),
       ('可展示', ['手機＋iPad 雙屏實機運作', '鏡頭前關閉行動數據，模型照樣措辭', '一鍵自動劇本，四個情境 45 秒跑完'])]
for i, (t, items) in enumerate(adv):
    cx = 0.7 + i * 4.1
    rect(s, cx, 2.35, 3.8, 3.75, CARD, LINE)
    text(s, cx + 0.25, 2.55, 3.4, 0.6, t, 24, ACC, True)
    text(s, cx + 0.25, 3.3, 3.35, 2.7, [f'・{x}' for x in items], 14, BODY, spacing=1.35)

# ── 16 預期效益 ───────────────────────────────────────
s = slide('⑦ 預期效益與落地評估', '預期效益', '使用者、車廠與技術三個層面', section='落地評估', page=16)
quad = [('使用者效益', '少被打斷：多數情況不出聲\n需要知道的事一句說清楚\n安全警示即時、方向直覺'),
        ('商業價值', '可整合進座艙域控制器（CDC）作為 HMI 決策層\n降低對雲端服務的依賴與成本\n每句話可追溯來源，便於驗證與責任釐清'),
        ('技術可行性', '1.5B 模型已在手機 GPU 實際執行\n規則可窮舉測試、檢查器可擴充\n事件格式穩定，雙屏已實機連線'),
        ('市場應用', '乘用車座艙 HMI\n商用車駕駛艙（長時間駕駛、訊號不穩路段）\n展示平台：多螢互動 Demo')]
for i, (t, d) in enumerate(quad):
    cx = 0.7 + (i % 2) * 6.05; cy = 2.3 + (i // 2) * 2.15
    rect(s, cx, cy, 5.8, 1.95, CARD, LINE)
    text(s, cx + 0.25, cy + 0.15, 5.3, 0.5, t, 18, ACC, True)
    text(s, cx + 0.25, cy + 0.65, 5.4, 1.25, d.split('\n'), 13.5, BODY, spacing=1.3)

# ── 17 落地路線與限制 ─────────────────────────────────
s = slide('⑦ 預期效益與落地評估', '落地路線與目前限制', '初審已完成的、決賽前要補的、之後要做的', section='落地評估', page=17)
cols = [('已完成（初審）', ACC, ['雙屏實機：儀表＋中控＋語音', '規則引擎、端側措辭、檢查器、知識庫', '斷網實測、自動化測試 16 項', '報名影片（實機錄影）']),
        ('決賽前（10/20）', NOTE, ['雲端增強：Jev 分流＋雲端聯網問答', '現場展示斷網與恢復連網', 'App 用詞修正（盲點標示為模擬）', '平板一機模式現場彩排']),
        ('決賽後', LO, ['接真實感測：CAN、CO2 感測器、盲點攝影機', '縮短措辭延遲、模型精簡（蒸餾）', '移植 AAOS／CDC、改用車上 Ethernet', '駕駛訪談與實車測試'])]
for i, (t, c, items) in enumerate(cols):
    cx = 0.7 + i * 4.1
    rect(s, cx, 2.3, 3.8, 0.06, c, None)
    text(s, cx, 2.5, 3.8, 0.5, t, 19, c, True)
    text(s, cx, 3.1, 3.8, 2.3, [f'・{x}' for x in items], 13.5, BODY, spacing=1.4)
rect(s, 0.7, 5.25, 11.9, 1.2, CARD, LINE)
text(s, 0.95, 5.33, 11.5, 1.1, [[('目前限制（主動揭露）：', NOTE, True)],
     '盲點與 CO2 為模擬注入事件，尚未接攝影機與感測器；主動告知措辭需 4.6–6.8 秒；雲端增強尚未上線；',
     'Demo 以手機兼任 AI Box、iPad 以網頁擔任中控（決賽改用 Android 平板一機模式）。'], 12.5, BODY)

# ── 18 封底 ───────────────────────────────────────────
s = slide(page=None)
text(s, 0.7, 2.3, 12, 1.2, '該說才說', 54, HI, True, PP_ALIGN.CENTER)
text(s, 0.7, 3.45, 12, 0.7, '會自己動手的安靜座艙 AI', 26, HI, False, PP_ALIGN.CENTER)
text(s, 0.7, 4.35, 12, 0.5, '規則決策｜端側 Qwen2.5-1.5B 措辭｜本地知識庫問答｜斷網照常運作（Offline First）', 16, ACC, False, PP_ALIGN.CENTER)
text(s, 0.7, 5.4, 12, 0.5, '方案說明影片：YouTube（不公開連結，見報名信件內文）', 14, LO, False, PP_ALIGN.CENTER)
text(s, 0.7, 5.85, 12, 0.5, '感知有度 Sense in Balance｜葉士愷・黃昱羲・胡馨', 14, LO, False, PP_ALIGN.CENTER)

out = r'F:\tmp\slides\out'
os.makedirs(out, exist_ok=True)
prs.save(out + r'\deck.pptx')
print('slides', len(prs.slides))
