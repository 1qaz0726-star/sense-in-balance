# 更新報名表與揭露表填寫版 → 另存新檔（不覆蓋原填寫版），再由 Word 轉 PDF
from docx import Document
import copy, re
D = r'C:\Users\葉士愷\Documents\友達智慧座艙\★送件用_填寫版'
TEAM = '感知有度 Sense in Balance'
PROP = '該說才說：會自己動手的安靜座艙 AI'

def set_par_text(p, text):
    runs = p.runs
    if not runs:
        p.add_run(text); return
    runs[0].text = text
    for r in runs[1:]: r.text = ''

def set_cell(cell, lines):
    """保留第一段格式，依行重建段落"""
    ps = cell.paragraphs
    first = ps[0]
    for p in ps[1:]:
        p._element.getparent().remove(p._element)
    set_par_text(first, lines[0])
    for ln in lines[1:]:
        np_ = copy.deepcopy(first._element)
        first._element.getparent().append(np_)
    for p, ln in zip(cell.paragraphs, lines):
        set_par_text(p, ln)

# ── 報名表 ───────────────────────────────────────────
SUMMARY = ('車上的提示越來越多，很多事系統自己就能處理，卻仍丟給駕駛。本提案「該說才說」讓座艙 AI 先動手、少開口：'
           '規則引擎依座艙狀態決定做什麼、多緊急、有沒有代價；能自動處理的事直接執行（例如車內 CO2 升高時自動切外循環），'
           '只有在有代價或有危險時才用一句話告知。安全警示由規則直接發出，不等模型；主動告知由手機端離線執行的 Qwen2.5-1.5B 措辭，'
           '並經檢查器把關字數與內容；停車時的提問查車上的知識庫，查不到就說查不到。數位儀表與中控兩屏依緊急度分工，'
           '斷網照常運作，決賽前加入雲端增強。')
n = len(re.sub(r'\s', '', SUMMARY))
assert n <= 300, n
d = Document(D + r'\附件二：報名表_填寫版.docx')
for p in d.paragraphs:
    t = p.text
    if t.startswith('團隊名稱：'): set_par_text(p, f'團隊名稱：{TEAM}')
    elif t.startswith('提案名稱：'): set_par_text(p, f'提案名稱：{PROP}')
    elif t.startswith('提案摘要'): set_par_text(p, f'提案摘要（300字內）：{SUMMARY}')
for r in d.tables[0].rows:
    if r.cells[0].text.strip() == '胡馨': set_cell(r.cells[3], ['碩士班四年級'])
d.save(D + r'\附件二：報名表_感知有度.docx')
print('摘要字數', n)

# ── 揭露表 ───────────────────────────────────────────
d = Document(D + r'\附件三：AI技術與第三方元件揭露表_填寫版.docx')
t = d.tables[0]
def row(label):
    for r in t.rows:
        if r.cells[0].text.strip().startswith(label): return r.cells[-1]
    raise KeyError(label)

set_cell(row('團隊名稱'), [TEAM])
set_cell(row('提案名稱'), [PROP])
set_cell(row('AI架構類型'), [
    '□ Cloud AI　■ Offline AI　□ Hybrid AI',
    '□ 概念規劃，尚未確定（創意組可勾選）',
    '※ 依送件時實際完成內容申報 Offline AI：決策、措辭、問答、語音皆於裝置端離線執行（報名影片含斷網實測）。'
    '雲端增強（Jev 意圖分流＋雲端聯網問答）預計於決賽前完成，完成並可現場展示斷網／復網後再更新為 Hybrid AI。'])
set_cell(row('AI Model名稱及版本'), [
    '本地端語言模型：Qwen2.5-1.5B-Instruct，q8 量化之 .litertlm 格式（LiteRT），於 Android 手機（Samsung Galaxy S21）端以 GPU 執行。',
    '用途：將規則引擎已決策完成之結構化狀態轉為一句繁體中文提示（主動告知 ≤15 字）。不做決策、不回答開放式問題；輸出未通過檢查即改用固定模板。',
    '語音：Android 系統文字轉語音（TTS，zh-TW），由裝置內建之語音引擎播放。',
    '雲端模型：無（現況）。規劃中：TypeSafe Jev（jev-1.13）用於提問意圖分流；Google Gemini 3.1 Flash Lite 經 OpenRouter 用於即時資訊問答。'])
set_cell(row('AI Model／Service Provider'), [
    'Qwen2.5-1.5B-Instruct：Alibaba Cloud（Qwen 團隊）；.litertlm 檔自 Hugging Face「litert-community」公開釋出取得。',
    'Android TTS：裝置內建語音引擎（Google 文字轉語音）。',
    '規劃中（尚未使用）：TypeSafe AI（Jev）、OpenRouter（Google Gemini）。'])
set_cell(row('使用之Framework'), [
    '1. Flutter 3.35.2／Dart 3.9：開發 Android HMI App（數位儀表、中控、網頁版中控）。',
    '2. flutter_gemma 0.12.6：於 Android 端載入並執行 .litertlm 模型（LiteRT）。',
    '3. flutter_tts 4.2.5：呼叫 Android 系統 TTS 播放語音。',
    '4. web_socket_channel 3.0.3 與 dart:io WebSocket：AI Box 與兩屏之事件同步。',
    '5. flutter_map 8.3.2、latlong2 0.10.1：中控地圖顯示。',
    '6. path_provider 2.1.5、file_picker 8.3.7：模型檔存放與匯入。',
    '7. 開發與驗證工具（不隨作品散布）：Ollama（開發期模型比較）、OpenCC（繁簡檢核）、Python。'])
set_cell(row('雲端AI服務或第三方API'), [
    '雲端 AI 服務：無（現況）。',
    '第三方資料服務：OpenStreetMap 圖磚伺服器（tile.openstreetmap.org），僅用於中控地圖底圖顯示。',
    '模型下載：Hugging Face（首次安裝時下載模型檔）。',
    '規劃中（決賽前）：TypeSafe Jev API、OpenRouter（Google Gemini 3.1 Flash Lite＋網路搜尋）；API Key 僅存於操作裝置之 App 設定，不寫入 APK。'])
set_cell(row('開源授權或商業授權類型'), [
    'Qwen2.5-1.5B-Instruct：Apache License 2.0',
    'Flutter：BSD-3-Clause　｜　flutter_gemma：MIT　｜　flutter_tts：MIT　｜　file_picker：MIT',
    'web_socket_channel、path_provider、flutter_map：BSD-3-Clause　｜　latlong2：Apache License 2.0',
    'OpenStreetMap 地圖資料：ODbL 1.0（© OpenStreetMap contributors），圖磚依 OSMF 使用政策使用',
    'Android 系統 TTS：裝置內建，隨裝置授權使用',
    '開發工具：Ollama：MIT　｜　OpenCC：Apache License 2.0'])
set_cell(row('使用之資料集及來源'), [
    '1. 團隊自產之模擬座艙狀態（CO2、車內外溫度、內外循環、乘員、車速）：由團隊撰寫之狀態產生器以程式合成，規則引擎以 12,000 組窮舉狀態驗證；非真實車輛資料。',
    '2. 盲點事件（方向、來車類型）為模擬注入之事件，未使用影像資料或攝影機。',
    '3. 本地知識庫 9 則問答：由團隊撰寫。',
    '※ 本提案未使用任何真實車聯網或其他競賽主辦單位提供之資料。'])
set_cell(row('使用之圖片、影片、音樂或其他素材來源'), [
    '1. 報名影片與簡報中之畫面：團隊以手機實際拍攝 App 運作（Samsung Galaxy S21、iPad），畫面中之第三方商標已做馬賽克。',
    '2. 旁白：團隊成員本人錄製；車內語音：Android 系統 TTS。',
    '3. 字型：Noto Sans TC（SIL Open Font License 1.1，影片字卡與字幕）；微軟正黑體（Windows 內建，簡報）。',
    '4. 地圖底圖：© OpenStreetMap contributors。',
    '5. 團隊照片：成員本人提供。',
    '6. 音樂：未使用。',
    '7. 影片製作工具：FFmpeg、Python（Pillow、OpenCV）、Remotion 4.0.534＋React（片頭、架構圖等動畫；Remotion License，個人及 3 人以下團隊免費）；字幕對時以語音辨識（Groq Whisper）檢查。'])
set_cell(row('既有技術、模型或元件'), [
    '使用下列第三方公開釋出之既有模型與元件：Qwen2.5-1.5B-Instruct、Flutter、flutter_gemma、flutter_tts、web_socket_channel、flutter_map、latlong2、path_provider、file_picker、Android 系統 TTS、OpenStreetMap 圖磚。',
    '上列模型均未經本團隊修改權重或重新訓練。',
    '本作品之研究原型（狀態產生器 states.py、事件介面 INTERFACE.md）由團隊於 2026 年 9 月為本競賽自行開發。'])
set_cell(row('本次競賽新增開發內容'), [
    '1. 規則引擎：依座艙狀態判定做什麼、緊急度（安全／主動告知／回答／不說話）與代價；與研究原型 12,000 組狀態逐筆一致。',
    '2. 執行層：可自動處理之事由系統直接執行（例：自動切外循環），有代價才告知。',
    '3. 端側措辭與檢查器：模型只負責措辭；檢查字數、簡體字、編造數字、誇大用語與「已完成動作被講成指令」，不通過即改用模板。',
    '4. 本地知識庫問答：停車時提問，查無資料即回覆「這個我查不到，請參考車主手冊」。',
    '5. 雙屏 HMI：數位儀表（駕駛最小資訊集合、盲點圖示）與中控（告知卡片、問答、地圖），含 iPad 網頁版中控；WebSocket 事件同步（1 秒心跳、3 秒逾時顯示失效）。',
    '6. 語音播報：與畫面同一句話，安全級可插話。',
    '7. 自動化測試 16 項（規則逐筆對照、檢查器、雙屏連線、網頁中控）。'])
set_cell(row('其他授權限制或注意事項'), [
    '1. 模型能力邊界（誠實揭露）：1.5B 模型於開放式車輛知識問答會產生幻覺，因此模型僅負責措辭；開放式問答一律查本地知識庫。',
    '2. 主動告知之措辭於手機端實測需 4.6–6.8 秒；安全警示不經模型，不受此延遲影響。',
    '3. 盲點與 CO2 事件為模擬注入，尚未連接真實感測器或攝影機。',
    '4. 端側模型採 q8 量化。'])
d.save(D + r'\附件三：AI技術與第三方元件揭露表_感知有度.docx')
print('ok')
